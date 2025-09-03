import argparse
from pathlib import Path
from typing import Optional

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet
from openpyxl.utils import get_column_letter


def _find_column_index(sheet: Worksheet, header_text: str, header_row: int = 1) -> int:
    """Find the 1-based column index by header value.

    Matching is done by exact string match after stripping whitespace.
    Raises ValueError if not found.
    """
    normalized_target = str(header_text).strip()
    for cell in sheet[header_row]:
        value = cell.value
        if value is None:
            continue
        if str(value).strip() == normalized_target:
            return cell.col_idx
    raise ValueError(f"Header '{header_text}' not found in row {header_row} of sheet '{sheet.title}'.")


def unmerge_and_fill_column(
    input_path: Path,
    output_path: Optional[Path] = None,
    sheet_name: Optional[str] = None,
    header_text: str = "单位名称",
    header_row: int = 1,
) -> Path:
    """Unmerge merged cells in a specific column and fill each row with the value.

    Only merged ranges that are entirely within the target column are unmerged.
    After unmerging, any blank cells in the column below the header are forward-filled.
    """
    if output_path is None:
        output_path = input_path.with_name(f"{input_path.stem}_unmerged{input_path.suffix}")

    wb = load_workbook(filename=str(input_path))
    ws = wb[sheet_name] if sheet_name else wb.active

    target_col_idx = _find_column_index(ws, header_text, header_row)
    target_col_letter = get_column_letter(target_col_idx)

    # 1) Unmerge merged cells that are vertical ranges within the target column
    # Use a copy of ranges since we'll mutate merged_cells during iteration
    ranges_to_process = list(ws.merged_cells.ranges)
    for merged_range in ranges_to_process:
        min_col, min_row, max_col, max_row = merged_range.bounds
        # Only unmerge if the range is fully within the target column (vertical merge)
        if min_col == max_col == target_col_idx:
            top_left_value = ws.cell(row=min_row, column=min_col).value
            ws.unmerge_cells(str(merged_range))
            for row in range(min_row, max_row + 1):
                ws.cell(row=row, column=target_col_idx).value = top_left_value

    # 2) Forward-fill any remaining blanks in the target column
    current_value = None
    for row in range(header_row + 1, ws.max_row + 1):
        cell = ws.cell(row=row, column=target_col_idx)
        if cell.value not in (None, ""):
            current_value = cell.value
        else:
            if current_value not in (None, ""):
                cell.value = current_value

    # Save results
    wb.save(str(output_path))
    return output_path


def main() -> None:
    parser = argparse.ArgumentParser(
        description=(
            "Unmerge merged cells in the '单位名称' column and fill each row with its value."
        )
    )
    parser.add_argument(
        "-i",
        "--input",
        type=Path,
        required=True,
        default="data/工作簿1.xlsx",
        help="Path to the input Excel file (.xlsx/.xlsm).",
    )
    parser.add_argument(
        "-o",
        "--output",
        type=Path,
        required=False,
        help="Optional output path; defaults to <name>_unmerged<ext> in the same folder.",
    )
    parser.add_argument(
        "-s",
        "--sheet",
        type=str,
        required=False,
        help="Sheet name to process; defaults to the active sheet if omitted.",
    )
    parser.add_argument(
        "--header-row",
        type=int,
        default=1,
        help="Header row index (1-based). Default: 1",
    )
    parser.add_argument(
        "--header-text",
        type=str,
        default="单位名称",
        help="Header text to locate the target column. Default: 单位名称",
    )

    args = parser.parse_args()

    output = unmerge_and_fill_column(
        input_path=args.input,
        output_path=args.output,
        sheet_name=args.sheet,
        header_text=args.header_text,
        header_row=args.header_row,
    )

    print(f"Saved: {output}")


if __name__ == "__main__":
    main()


